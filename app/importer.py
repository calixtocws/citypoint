import re
from openpyxl import load_workbook
from . import db
from .pools import refresh_statuses

def clean(v): return "" if v is None else str(v).strip()
def booth_group_key(v):
    s=clean(v)
    if not s: return ""
    parts=[p.strip() for p in s.split(',') if p.strip()]
    return ', '.join(parts) if len(parts)>1 else s
def norm_token(tok,prefix=''):
    tok=clean(tok).replace(' ','')
    tok=re.sub(r'^([A-Za-z])-([0-9]+)$', r'\1\2', tok)
    m=re.match(r'^([A-Za-z]+)([0-9]+)$', tok)
    if m: return m.group(1).upper()+m.group(2), m.group(1).upper()
    if re.match(r'^[0-9]+$', tok) and prefix: return prefix+tok, prefix
    return tok.upper(), prefix
def booth_from_location(location):
    location=clean(location)
    if 'Booth' not in location: return ''
    tmp=re.sub(r'^.*?Booths?-','',location)
    out=[]; prefix=''
    for p in [x.strip() for x in tmp.split('-') if x.strip()]:
        b,prefix=norm_token(p,prefix)
        if b: out.append(b)
    return ', '.join(out)
def parse_vlans(v):
    s=clean(v)
    if not s or s.upper().startswith('DMH-'): return []
    if re.match(r'^\d{6}$',s): return [int(s[:3]),int(s[3:])]
    if re.match(r'^\d{3}-\d{3}$',s):
        a,b=map(int,s.split('-')); return list(range(a,b+1))
    if re.match(r'^\d+(\.0)?$',s): return [int(float(s))]
    return []
def import_excel(path):
    wb=load_workbook(path,data_only=True)
    db.execute("DELETE FROM booths"); db.execute("DELETE FROM customers")
    imported={"booths":0,"customers":0,"subnet_links":0}
    subnet_by_vlan={}; subnet_by_booth={}
    if 'cust vs subnet' in wb.sheetnames:
        ws=wb['cust vs subnet']
        for r in range(1,ws.max_row+1):
            cust=clean(ws.cell(r,1).value); cidr=clean(ws.cell(r,2).value); gw=clean(ws.cell(r,3).value); loc=clean(ws.cell(r,4).value); vlan=clean(ws.cell(r,7).value)
            booth=booth_from_location(loc)
            if vlan and cidr:
                try: subnet_by_vlan[int(float(vlan))]={"cidr":cidr,"gateway":gw,"customer":cust,"location":loc}
                except Exception: pass
            if booth and cidr: subnet_by_booth[booth]={"cidr":cidr,"gateway":gw,"customer":cust,"location":loc,"vlan":vlan}
    if 'vlan vs booth' not in wb.sheetnames: return {"error":"Missing sheet: vlan vs booth"}
    ws=wb['vlan vs booth']
    for r in range(2,ws.max_row+1):
        booth=booth_group_key(ws.cell(r,1).value); licensee=clean(ws.cell(r,2).value); legal=clean(ws.cell(r,3).value); vlan_raw=clean(ws.cell(r,4).value); vlans=parse_vlans(vlan_raw)
        if booth:
            db.execute("INSERT OR IGNORE INTO booths(booth_group,status) VALUES (?, 'available')", (booth,)); imported['booths']+=1
        if not licensee or licensee.upper()=='VACANT': continue
        chosen_vlan=None; notes=["Imported from original workbook"]
        if len(vlans)==1: chosen_vlan=vlans[0]
        elif len(vlans)>1: notes.append(f"Original workbook has multiple VLANs: {', '.join(map(str,vlans))}. VLAN left empty for review.")
        elif vlan_raw: notes.append(f"Original workbook VLAN value not parsed: {vlan_raw}")
        subnet_cidr=None
        if chosen_vlan and chosen_vlan in subnet_by_vlan:
            subnet_cidr=subnet_by_vlan[chosen_vlan]['cidr']; imported['subnet_links']+=1
        elif booth in subnet_by_booth:
            subnet_cidr=subnet_by_booth[booth]['cidr']; imported['subnet_links']+=1
        db.execute("INSERT INTO customers(licensee,legal_name,booth_group,vlan_id,subnet_cidr,notes) VALUES (?, ?, ?, ?, ?, ?)", (licensee,legal,booth,chosen_vlan,subnet_cidr,'; '.join(notes)))
        imported['customers']+=1
    refresh_statuses(); return imported
